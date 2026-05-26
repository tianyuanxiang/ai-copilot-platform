import base64
import math
import re
import zipfile
from io import BytesIO
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import httpx

try:
    from docx import Document
except ImportError:  # pragma: no cover
    Document = None

try:
    from langchain_text_splitters import RecursiveCharacterTextSplitter
except ImportError:  # pragma: no cover
    RecursiveCharacterTextSplitter = None

from app.core.config import Settings, get_settings
from app.schemas.knowledge import (
    ChildChunk,
    EmbedResponse,
    ParentChunk,
    ParseResponse,
    RerankResponse,
    RetrievedChunk,
)

DASHSCOPE_EMBEDDINGS_URL = "https://dashscope.aliyuncs.com/compatible-mode/v1/embeddings"
DASHSCOPE_RERANK_URL = "https://dashscope.aliyuncs.com/api/v1/services/rerank/text-rerank/text-rerank"

MAX_EMBEDDING_TOKENS = 8192
PARENT_TARGET_CHARS = 1000
PARENT_MAX_CHARS = 1200
CHILD_TARGET_CHARS = 350
CHILD_MAX_CHARS = 420
CHILD_OVERLAP_CHARS = 50
WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
TEXT_FORMAT_MARKS = "\ufeff\u200b\u200c\u200d"
SUPPORTED_EMBEDDING_INPUT_TYPES = {"document", "query"}


class ParseInputError(ValueError):
    pass


class EmbeddingInputError(ValueError):
    pass


class EmbeddingProviderError(RuntimeError):
    pass


async def parse_document(file_name: str, content: str, file_type: str = "") -> ParseResponse:
    document_type = _detect_document_type(file_name, file_type)
    if document_type == "docx":
        text = _extract_docx_text(_decode_docx_content(content))
        parser = "docx-structured"
    elif document_type in {"md", "markdown"}:
        text = _normalize_markdown_text(content)
        parser = "markdown-structure"
    elif document_type == "txt":
        text = _normalize_text(content)
        parser = "plain-text"
    else:
        text = _normalize_text(content)
        parser = "plain-text"

    parents = chunk_document(text, document_type)
    return ParseResponse(
        text=text,
        metadata={
            "file_name": file_name,
            "file_type": document_type,
            "parser": parser,
            "parent_count": len(parents),
            "child_count": sum(len(parent.children) for parent in parents),
            "splitter": "langchain-recursive" if RecursiveCharacterTextSplitter else "fallback-recursive",
        },
        parents=parents,
    )


def chunk_document(text: str, document_type: str = "txt") -> list[ParentChunk]:
    normalized = _normalize_text(text)
    if not normalized:
        return []
    if document_type in {"md", "markdown"}:
        parent_texts = _split_markdown_parent_texts(normalized)
    else:
        parent_texts = _split_text(normalized, PARENT_MAX_CHARS, 0)
    return _build_parent_chunks(parent_texts)


def _build_parent_chunks(parent_texts: list[str]) -> list[ParentChunk]:
    parents: list[ParentChunk] = []
    for parent_index, parent_text in enumerate(parent_texts):
        child_texts = _split_child_chunks(parent_text)
        parents.append(
            ParentChunk(
                parent_index=parent_index,
                content=parent_text,
                token_count=_estimate_chunk_tokens(parent_text),
                children=[
                    ChildChunk(
                        chunk_index=child_index,
                        content=child_text,
                        token_count=_estimate_chunk_tokens(child_text),
                    )
                    for child_index, child_text in enumerate(child_texts)
                ],
            )
        )
    return parents


def _split_markdown_parent_texts(text: str) -> list[str]:
    sections = _markdown_to_retrieval_sections(text)
    parent_texts: list[str] = []
    for section in sections:
        parent_texts.extend(_split_text(section, PARENT_MAX_CHARS, 0))
    return [item for item in parent_texts if item.strip()]


def _markdown_to_retrieval_sections(text: str) -> list[str]:
    """Convert Markdown into readable retrieval sections before chunking."""

    sections: list[str] = []
    heading_stack: list[str] = []
    current: list[str] = []
    in_code_block = False

    def flush() -> None:
        body = _normalize_text("\n".join(item for item in current if item.strip()))
        if body:
            prefix = " > ".join(heading_stack)
            if prefix and not body.startswith(prefix):
                body = f"{prefix}\n\n{body}"
            sections.append(_normalize_text(body))
        current.clear()

    for raw_line in _normalize_text(text).split("\n"):
        stripped = raw_line.strip()
        if stripped.startswith("```") or stripped.startswith("~~~"):
            in_code_block = not in_code_block
            continue
        if not in_code_block and _is_markdown_table_separator(stripped):
            continue
        if not in_code_block and (stripped.upper() == "[TOC]" or re.fullmatch(r"-{3,}", stripped)):
            continue

        heading = None if in_code_block else re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading:
            flush()
            level = len(heading.group(1))
            title = _clean_inline_markdown(heading.group(2))
            heading_stack[:] = heading_stack[: level - 1]
            heading_stack.append(title)
            current.append(title)
            continue

        current.append(stripped if in_code_block else _clean_markdown_line(raw_line))

    flush()
    return sections or [_normalize_markdown_text(text)]


def _split_child_chunks(parent_text: str) -> list[str]:
    return _split_text(parent_text, CHILD_MAX_CHARS, CHILD_OVERLAP_CHARS)


def _split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    clean_text = _normalize_text(text)
    if not clean_text:
        return []
    if RecursiveCharacterTextSplitter is not None:
        splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=overlap,
            separators=["\n\n", "\n", "。", "！", "？", "；", ";", ".", "!", "?", "，", ",", " ", ""],
            keep_separator="end",
        )
        return [_trim_fragment(item) for item in splitter.split_text(clean_text) if _trim_fragment(item)]
    return _fallback_split(clean_text, chunk_size, overlap)


def _fallback_split(text: str, chunk_size: int, overlap: int) -> list[str]:
    if len(text) <= chunk_size:
        return [text]
    sentences = re.split(r"(?<=[。！？；;.!?])\s+", text)
    chunks: list[str] = []
    current = ""
    for sentence in sentences:
        if not sentence:
            continue
        candidate = f"{current} {sentence}".strip() if current else sentence
        if len(candidate) > chunk_size and current:
            chunks.append(current)
            prefix = current[-overlap:].strip() if overlap > 0 else ""
            current = f"{prefix} {sentence}".strip() if prefix else sentence
        elif len(sentence) > chunk_size:
            chunks.extend(_window_split(sentence, chunk_size, overlap))
            current = ""
        else:
            current = candidate
    if current:
        chunks.append(current)
    return [_trim_fragment(chunk) for chunk in chunks if _trim_fragment(chunk)]


def _window_split(text: str, size: int, overlap: int) -> list[str]:
    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + size)
        chunks.append(text[start:end])
        if end == len(text):
            break
        start = max(end - overlap, start + 1)
    return chunks


def _trim_fragment(text: str) -> str:
    text = _normalize_text(text)
    text = re.sub(r"^[，,。；;：:\s]+", "", text)
    return text.strip()


def _detect_document_type(file_name: str, file_type: str) -> str:
    clean_type = file_type.strip().lower().lstrip(".")
    if clean_type:
        return clean_type
    suffix = Path(file_name).suffix.lower().lstrip(".")
    return suffix or "txt"


def _normalize_text(text: str) -> str:
    normalized = text.replace("\r\n", "\n").replace("\r", "\n")
    normalized = normalized.lstrip(TEXT_FORMAT_MARKS)
    normalized = re.sub(rf"\n[{TEXT_FORMAT_MARKS}]+", "\n", normalized)
    normalized = re.sub(r"[ \t]+\n", "\n", normalized)
    normalized = re.sub(r"\n{3,}", "\n\n", normalized)
    return normalized.strip()


def _normalize_markdown_text(text: str) -> str:
    return _normalize_text("\n".join(_markdown_to_retrieval_sections(text)))


def _clean_markdown_line(line: str) -> str:
    stripped = line.strip()
    if not stripped:
        return ""
    stripped = re.sub(r"^\s{0,3}>\s?", "", stripped)
    stripped = re.sub(r"^\s*[-*+]\s+", "", stripped)
    stripped = re.sub(r"^\s*\d+[.)]\s+", "", stripped)
    if "|" in stripped:
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        stripped = "；".join(cell for cell in cells if cell)
    return _clean_inline_markdown(stripped)


def _clean_inline_markdown(text: str) -> str:
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"(\*\*|__)(.*?)\1", r"\2", text)
    text = re.sub(r"(\*|_)(.*?)\1", r"\2", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def _is_markdown_table_separator(line: str) -> bool:
    if "|" not in line:
        return False
    cells = [cell.strip() for cell in line.strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell or "") for cell in cells)


def _decode_docx_content(content: str) -> bytes:
    raw = content.strip()
    if raw.startswith("data:"):
        _, _, raw = raw.partition(",")
    try:
        data = base64.b64decode(raw, validate=True)
    except ValueError as exc:
        raise ParseInputError("docx content must be base64 encoded") from exc
    if not zipfile.is_zipfile(BytesIO(data)):
        raise ParseInputError("docx content is not a valid .docx zip package")
    return data


def _extract_docx_text(data: bytes) -> str:
    if Document is not None:
        try:
            document = Document(BytesIO(data))
            parts: list[str] = []
            parts.extend(paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip())
            for table in document.tables:
                rows: list[str] = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        rows.append(" | ".join(cells))
                if rows:
                    parts.append("\n".join(rows))
            return _normalize_text("\n\n".join(parts))
        except Exception:
            return _extract_docx_text_with_xml(data)
    return _extract_docx_text_with_xml(data)


def _extract_docx_text_with_xml(data: bytes) -> str:
    try:
        with zipfile.ZipFile(BytesIO(data)) as package:
            document_xml = package.read("word/document.xml")
    except KeyError as exc:
        raise ParseInputError("docx is missing word/document.xml") from exc
    except zipfile.BadZipFile as exc:
        raise ParseInputError("docx content is not a readable zip package") from exc

    root = ElementTree.fromstring(document_xml)
    body = root.find(f"{WORD_NAMESPACE}body")
    if body is None:
        return ""

    parts: list[str] = []
    for child in body:
        if child.tag == f"{WORD_NAMESPACE}p":
            paragraph = _docx_paragraph_text(child)
            if paragraph:
                parts.append(paragraph)
        elif child.tag == f"{WORD_NAMESPACE}tbl":
            table = _docx_table_text(child)
            if table:
                parts.append(table)
    return _normalize_text("\n\n".join(parts))


def _docx_paragraph_text(paragraph: ElementTree.Element) -> str:
    fragments: list[str] = []
    for node in paragraph.iter():
        if node.tag == f"{WORD_NAMESPACE}t" and node.text:
            fragments.append(node.text)
        elif node.tag == f"{WORD_NAMESPACE}tab":
            fragments.append("\t")
        elif node.tag in {f"{WORD_NAMESPACE}br", f"{WORD_NAMESPACE}cr"}:
            fragments.append("\n")
    return "".join(fragments).strip()


def _docx_table_text(table: ElementTree.Element) -> str:
    rows: list[str] = []
    for row in table.findall(f".//{WORD_NAMESPACE}tr"):
        cells: list[str] = []
        for cell in row.findall(f"{WORD_NAMESPACE}tc"):
            paragraphs = [_docx_paragraph_text(paragraph) for paragraph in cell.findall(f"{WORD_NAMESPACE}p")]
            cell_text = " ".join(item for item in paragraphs if item).strip()
            cells.append(cell_text)
        if any(cells):
            rows.append(" | ".join(cells))
    return "\n".join(rows).strip()


def _estimate_chunk_tokens(text: str) -> int:
    return max(1, math.ceil(len(text) / 4))


async def embed_texts(
    texts: list[str],
    input_type: str = "document",
    settings: Settings | None = None,
    transport: httpx.AsyncBaseTransport | None = None,
) -> EmbedResponse:
    settings = settings or get_settings()
    input_type = _normalize_embedding_input_type(input_type)
    normalized = _validate_embedding_texts(texts)

    if settings.embedding_provider == "dashscope":
        return await _embed_dashscope(normalized, input_type, settings, transport)

    token_counts = [_estimate_tokens(text) for text in normalized]
    return EmbedResponse(
        vectors=[_mock_vector(text, settings.embedding_dimension) for text in normalized],
        token_counts=token_counts,
        total_tokens=sum(token_counts),
        model=settings.embedding_model,
        dimension=settings.embedding_dimension,
        mode="mock-embedding",
    )


def _normalize_embedding_input_type(input_type: str) -> str:
    clean = (input_type or "document").strip().lower()
    if clean not in SUPPORTED_EMBEDDING_INPUT_TYPES:
        raise EmbeddingInputError("input_type must be document or query")
    return clean


def _validate_embedding_texts(texts: list[str]) -> list[str]:
    if not texts:
        raise EmbeddingInputError("texts must contain at least one item")
    normalized: list[str] = []
    for index, text in enumerate(texts):
        clean_text = text.strip()
        if not clean_text:
            raise EmbeddingInputError(f"texts[{index}] must not be blank")
        token_count = _estimate_tokens(clean_text)
        if token_count > MAX_EMBEDDING_TOKENS:
            raise EmbeddingInputError(f"texts[{index}] exceeds {MAX_EMBEDDING_TOKENS} tokens")
        normalized.append(clean_text)
    return normalized


def _estimate_tokens(text: str) -> int:
    return max(1, len(text.split()))


def _mock_vector(text: str, dimension: int) -> list[float]:
    seed = sum(ord(char) for char in text) or 1
    return [float(((seed + idx) % 997) / 997) for idx in range(dimension)]


async def _embed_dashscope(
    texts: list[str],
    input_type: str,
    settings: Settings,
    transport: httpx.AsyncBaseTransport | None,
) -> EmbedResponse:
    if not settings.embedding_api_key:
        raise EmbeddingProviderError("embedding api key is empty; set embedding.api_key or DASHSCOPE_API_KEY")

    batch_size = max(1, min(settings.embedding_batch_size, 10))
    vectors: list[list[float]] = []
    token_counts: list[int] = []
    total_tokens = 0

    async with httpx.AsyncClient(timeout=settings.embedding_timeout_seconds, transport=transport) as client:
        for start in range(0, len(texts), batch_size):
            batch = texts[start : start + batch_size]
            payload = {
                "model": settings.embedding_model,
                "input": batch,
                "dimensions": settings.embedding_dimension,
                "encoding_format": "float",
            }
            response = await client.post(
                DASHSCOPE_EMBEDDINGS_URL,
                headers={"Authorization": f"Bearer {settings.embedding_api_key}"},
                json=payload,
            )
            if response.status_code < 200 or response.status_code >= 300:
                raise EmbeddingProviderError(
                    f"dashscope embedding returned {response.status_code}: {response.text[:4096]}"
                )

            body = response.json()
            batch_vectors = _extract_dashscope_vectors(body, settings.embedding_dimension)
            if len(batch_vectors) != len(batch):
                raise EmbeddingProviderError(
                    f"dashscope embedding returned {len(batch_vectors)} vectors for {len(batch)} texts"
                )
            vectors.extend(batch_vectors)

            usage_total = _extract_usage_total(body)
            if usage_total > 0:
                batch_token_counts = _allocate_usage_tokens(usage_total, len(batch))
            else:
                batch_token_counts = [_estimate_tokens(text) for text in batch]
            token_counts.extend(batch_token_counts)
            total_tokens += sum(batch_token_counts)

    return EmbedResponse(
        vectors=vectors,
        token_counts=token_counts,
        total_tokens=total_tokens,
        model=settings.embedding_model,
        dimension=settings.embedding_dimension,
        mode="dashscope-openai-compatible-embedding",
    )


def _extract_dashscope_vectors(body: dict[str, Any], expected_dimension: int) -> list[list[float]]:
    data = body.get("data")
    if not isinstance(data, list):
        raise EmbeddingProviderError("dashscope embedding response missing data list")
    vectors: list[list[float]] = []
    for item in data:
        if not isinstance(item, dict):
            raise EmbeddingProviderError("dashscope embedding response contains invalid data item")
        embedding = item.get("embedding")
        if not isinstance(embedding, list):
            raise EmbeddingProviderError("dashscope embedding response missing embedding vector")
        if len(embedding) != expected_dimension:
            raise EmbeddingProviderError(
                f"dashscope embedding dimension is {len(embedding)}, expected {expected_dimension}"
            )
        vectors.append([float(value) for value in embedding])
    return vectors


def _extract_usage_total(body: dict[str, Any]) -> int:
    usage = body.get("usage")
    if not isinstance(usage, dict):
        return 0
    value = usage.get("total_tokens") or usage.get("prompt_tokens") or 0
    try:
        return int(value)
    except (TypeError, ValueError):
        return 0


def _allocate_usage_tokens(total_tokens: int, item_count: int) -> list[int]:
    base = total_tokens // item_count
    remainder = total_tokens % item_count
    return [base + (1 if index < remainder else 0) for index in range(item_count)]


async def rerank(query: str, chunks: list[RetrievedChunk], top_k: int) -> RerankResponse:
    settings = get_settings()
    if not chunks:
        return RerankResponse(chunks=[], mode="empty")
    top_k = max(1, min(top_k or settings.rerank_top_n, len(chunks)))
    if settings.rerank_provider == "dashscope" and settings.rerank_api_key:
        return await _rerank_dashscope(query, chunks, top_k, settings)
    return RerankResponse(chunks=_mock_rerank(query, chunks, top_k), mode="mock-rerank")


def _mock_rerank(query: str, chunks: list[RetrievedChunk], top_k: int) -> list[RetrievedChunk]:
    terms = _query_terms(query)

    def score(item: RetrievedChunk) -> float:
        content = item.content.lower()
        overlap = sum(1 for term in terms if term and term.lower() in content)
        return item.score + overlap

    return sorted(chunks, key=score, reverse=True)[:top_k]


async def _rerank_dashscope(query: str, chunks: list[RetrievedChunk], top_k: int, settings: Settings) -> RerankResponse:
    payload = {
        "model": settings.rerank_model,
        "input": {
            "query": query,
            "documents": [chunk.content for chunk in chunks],
        },
        "parameters": {
            "top_n": top_k,
            "return_documents": True,
        },
    }
    async with httpx.AsyncClient(timeout=settings.rerank_timeout_seconds) as client:
        response = await client.post(
            DASHSCOPE_RERANK_URL,
            headers={"Authorization": f"Bearer {settings.rerank_api_key}"},
            json=payload,
        )
    if response.status_code < 200 or response.status_code >= 300:
        raise EmbeddingProviderError(f"dashscope rerank returned {response.status_code}: {response.text[:4096]}")

    results = _extract_dashscope_rerank_results(response.json())
    ranked: list[RetrievedChunk] = []
    for item in results:
        index = item["index"]
        if index < 0 or index >= len(chunks):
            continue
        chunk = chunks[index].model_copy()
        chunk.score = item["score"]
        ranked.append(chunk)
    return RerankResponse(chunks=ranked[:top_k], mode="dashscope-rerank")


def _extract_dashscope_rerank_results(body: dict[str, Any]) -> list[dict[str, Any]]:
    output = body.get("output")
    if isinstance(output, dict):
        candidates = output.get("results") or output.get("documents") or []
    else:
        candidates = body.get("results") if isinstance(body.get("results"), list) else []

    results: list[dict[str, Any]] = []
    for item in candidates:
        if not isinstance(item, dict):
            continue
        index = item.get("index")
        score = item.get("relevance_score", item.get("score", 0))
        try:
            results.append({"index": int(index), "score": float(score)})
        except (TypeError, ValueError):
            continue
    return sorted(results, key=lambda item: item["score"], reverse=True)


def _query_terms(query: str) -> list[str]:
    terms = re.findall(r"[\w\u4e00-\u9fff]+", query)
    return [term for term in terms if len(term) > 1]
